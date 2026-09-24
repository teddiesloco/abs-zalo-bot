"""Dựng tệp Word / PDF / PowerPoint / Excel từ nội dung bot soạn, để gửi vào nhóm Zalo.

Vì sao cần: người trong nhóm không có công cụ ghi tệp hay chạy lệnh (mở ra là mở cả máy
chủ). Mô-đun này chỉ nhận *nội dung* — chữ Markdown, danh sách slide, bảng — rồi dựng tệp
trong thư mục tạm mà công cụ gọi truyền vào. Không đọc tệp nào trên máy, không tải ảnh,
không chạy lệnh; mọi kích thước đều có trần để một lời nhờ không làm treo bot.

Trình bày: thầy cô dùng tệp để in, chiếu, gửi tiếp — tệp trơn trông như bản nháp. Word theo
kỹ thuật trình bày Nghị định 30 (chữ đen, khổ và lề chuẩn — xem build_docx) để in nộp được
ngay; PDF, PowerPoint, Excel dùng chung bảng màu (xanh đậm + xanh nhấn + nền nhạt), bảng có
dòng tiêu đề nền đậm và dòng xen kẽ.
"""

from __future__ import annotations

import os
import re
import unicodedata
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

FORMATS = ("docx", "pptx", "xlsx", "pdf")
MAX_TEXT_CHARS = 30_000
MAX_SLIDES = 40
MAX_BULLETS_PER_SLIDE = 15
MAX_SHEETS = 5
MAX_ROWS = 2_000
MAX_COLS = 30
MAX_CELL_CHARS = 1_000
MAX_TITLE_CHARS = 200

# Bảng màu chung (hex không dấu #).
NAVY = "1F4E79"
ACCENT = "2E75B6"
LIGHT = "DEEAF6"
ZEBRA = "F2F7FC"
INK = "263238"
MUTED = "7A8A99"
GRID = "BFC9D4"

# Font có đủ dấu tiếng Việt cho PDF: (thường, đậm). Biến môi trường thắng. Arial không có
# chỉ số dưới (H₂O thành "HO") nên Windows dùng Segoe UI; DejaVu Sans trên Linux có đủ.
_PDF_FONTS = (
    (r"C:\Windows\Fonts\segoeui.ttf", r"C:\Windows\Fonts\segoeuib.ttf"),
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ("/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf", "/usr/share/fonts/truetype/noto/NotoSans-Bold.ttf"),
)
# Font dự phòng cho ký hiệu font chính thiếu (⇒ trên Segoe UI).
_PDF_FALLBACK_FONTS = (
    r"C:\Windows\Fonts\seguisym.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
)
# PowerPoint/Excel: Calibri có đủ chỉ số trên/dưới, Office nào cũng sẵn.
OFFICE_FONT = "Calibri"


class FileSpecError(ValueError):
    """Nội dung không dựng được tệp — thông điệp đưa thẳng cho người nhờ."""


def _rgb(hex_color: str) -> Tuple[int, int, int]:
    return tuple(int(hex_color[i:i + 2], 16) for i in (0, 2, 4))


# ---------------------------------------------------------------- nội dung


def safe_filename(name: str, fmt: str) -> str:
    """Tên tệp dễ đọc (giữ dấu cách, chữ tiếng Việt), bỏ ký tự có thể thành đường dẫn."""
    stem = unicodedata.normalize("NFC", str(name or "").strip())
    stem = re.sub(r"\.(docx|pptx|xlsx|pdf)$", "", stem, flags=re.IGNORECASE)
    stem = re.sub(r"[\\/:*?\"<>|\x00-\x1f]+", " ", stem)
    stem = re.sub(r"\s+", " ", stem).strip("._ ")[:80].strip("._ ")
    return f"{stem or 'Tài liệu'}.{fmt}"


def parse_blocks(markdown: str) -> List[Tuple]:
    """Tách Markdown đơn giản thành khối: tiêu đề, gạch đầu dòng, đánh số, bảng, đoạn."""
    blocks: List[Tuple] = []
    table: List[List[str]] = []

    def flush_table() -> None:
        if table:
            blocks.append(("table", [row[:] for row in table]))
            table.clear()

    for raw in str(markdown or "").splitlines():
        line = raw.rstrip()
        stripped = line.strip()
        if stripped.startswith("|") and stripped.endswith("|") and len(stripped) > 1:
            cells = [cell.strip() for cell in stripped[1:-1].split("|")]
            if not all(re.fullmatch(r":?-{2,}:?", cell) for cell in cells if cell):
                table.append(cells)
            continue
        flush_table()
        if not stripped:
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        bullet = re.match(r"^([-*•+])\s+(.*)$", stripped)
        number = re.match(r"^(\d+[.)])\s+(.*)$", stripped)
        indent = len(line) - len(line.lstrip())
        if heading:
            blocks.append(("heading", min(len(heading.group(1)), 3), heading.group(2).strip()))
        elif bullet:
            blocks.append(("bullet", 1 if indent >= 2 else 0, bullet.group(2).strip()))
        elif number:
            blocks.append(("number", 0, f"{number.group(1)} {number.group(2).strip()}"))
        else:
            blocks.append(("para", 0, stripped))
    flush_table()
    return blocks


def inline_runs(text: str) -> List[Tuple[str, bool, bool]]:
    """Chia một dòng thành các đoạn (chữ, đậm, nghiêng) theo **đậm** và *nghiêng*."""
    runs: List[Tuple[str, bool, bool]] = []
    for part in re.split(r"(\*\*[^*]+\*\*|(?<![\w*])\*[^*\s][^*]*\*(?![\w*]))", str(text)):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            runs.append((part[2:-2], True, False))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            runs.append((part[1:-1], False, True))
        else:
            runs.append((part, False, False))
    return runs


def _plain(text: str) -> str:
    return "".join(chunk for chunk, _b, _i in inline_runs(text))


def _square(rows: List[List[str]]) -> List[List[str]]:
    width = max(len(row) for row in rows)
    return [row + [""] * (width - len(row)) for row in rows]


def _check_text(title: str, content: str) -> None:
    if len(title) > MAX_TITLE_CHARS:
        raise FileSpecError(f"tiêu đề dài quá {MAX_TITLE_CHARS} ký tự")
    if not content.strip():
        raise FileSpecError("cần `content` — nội dung văn bản của tệp")
    if len(content) > MAX_TEXT_CHARS:
        raise FileSpecError(f"nội dung dài quá {MAX_TEXT_CHARS} ký tự — chia thành nhiều tệp nhỏ hơn")


def _normalize_slides(slides: Any) -> List[Dict[str, Any]]:
    if not isinstance(slides, list) or not slides:
        raise FileSpecError("cần `slides` — danh sách slide, mỗi slide có `title` và `bullets`")
    if len(slides) > MAX_SLIDES:
        raise FileSpecError(f"tối đa {MAX_SLIDES} slide mỗi tệp")
    result = []
    for index, slide in enumerate(slides, 1):
        if not isinstance(slide, dict):
            raise FileSpecError(f"slide {index} phải có dạng {{title, bullets}}")
        bullets = slide.get("bullets") or []
        if not isinstance(bullets, list):
            raise FileSpecError(f"`bullets` của slide {index} phải là danh sách")
        if len(bullets) > MAX_BULLETS_PER_SLIDE:
            raise FileSpecError(f"slide {index} có quá {MAX_BULLETS_PER_SLIDE} ý — tách thành nhiều slide")
        title = str(slide.get("title") or "").strip()[:MAX_TITLE_CHARS]
        items = [str(b).strip()[:MAX_CELL_CHARS] for b in bullets if str(b).strip()]
        if not title and not items:
            raise FileSpecError(f"slide {index} trống")
        result.append({"title": title, "bullets": items})
    return result


def _normalize_sheets(sheets: Any) -> List[Dict[str, Any]]:
    if not isinstance(sheets, list) or not sheets:
        raise FileSpecError("cần `sheets` — danh sách trang tính, mỗi trang có `name` và `rows`")
    if len(sheets) > MAX_SHEETS:
        raise FileSpecError(f"tối đa {MAX_SHEETS} trang tính mỗi tệp")
    result, used = [], set()
    for index, sheet in enumerate(sheets, 1):
        if not isinstance(sheet, dict) or not isinstance(sheet.get("rows"), list) or not sheet["rows"]:
            raise FileSpecError(f"trang tính {index} cần `rows` — danh sách các dòng")
        rows = sheet["rows"]
        if len(rows) > MAX_ROWS:
            raise FileSpecError(f"trang tính {index} quá {MAX_ROWS} dòng")
        clean_rows = []
        for row in rows:
            cells = row if isinstance(row, list) else [row]
            if len(cells) > MAX_COLS:
                raise FileSpecError(f"trang tính {index} quá {MAX_COLS} cột")
            clean_rows.append([_cell(value) for value in cells])
        name = re.sub(r"[\[\]:*?/\\]", " ", str(sheet.get("name") or f"Trang {index}")).strip()[:31] or f"Trang {index}"
        while name in used:
            name = f"{name[:28]} {index}"
        used.add(name)
        result.append({"name": name, "rows": clean_rows})
    return result


def _cell(value: Any) -> Any:
    if isinstance(value, bool) or value is None:
        return "" if value is None else str(value)
    if isinstance(value, (int, float)):
        return value
    text = str(value)[:MAX_CELL_CHARS]
    # Chữ bắt đầu bằng "=" là công thức Excel: người lạ không được cài công thức vào tệp.
    return "'" + text if text.startswith(("=", "+", "-", "@")) and not re.fullmatch(r"[+-]?\d+([.,]\d+)?", text) else text


# ---------------------------------------------------------------- Word


def _docx_fonts(target, name: str) -> None:
    """Đặt font cho style/run, bỏ font theo theme để Word không tự đổi về Calibri."""
    from docx.oxml.ns import qn

    target.font.name = name
    rpr = target.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attr), None)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)


def _docx_row_flags(row, header: bool) -> None:
    """Hàng tiêu đề lặp lại mỗi trang; không hàng nào bị tách qua hai trang."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tr_pr = row._tr.get_or_add_trPr()
    for tag in (("w:cantSplit", "w:tblHeader") if header else ("w:cantSplit",)):
        flag = OxmlElement(tag)
        flag.set(qn("w:val"), "true")
        tr_pr.append(flag)


def _docx_page_number(paragraph) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        run = paragraph.add_run()
        if kind:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        else:
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = text
        run._r.append(element)


def build_docx(title: str, content: str, path: Path) -> None:
    """Word theo kỹ thuật trình bày Nghị định 30/2020/NĐ-CP (skill soan-van-ban-doan, profile nd30).

    A4; lề trên/dưới 20 mm, trái 30 mm, phải 15 mm; Times New Roman 14 màu đen; căn đều, thụt
    dòng đầu 1 cm, cách đoạn 6 pt; gạch đầu dòng gõ tay (không dùng danh sách tự động của
    Word); bảng lặp hàng tiêu đề, không tách hàng; số trang giữa lề trên, không hiện ở trang 1.
    Giáo án, đề, danh sách không phải văn bản hành chính nên không có khối quốc hiệu.
    """
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    body_font, size = "Times New Roman", 14
    black = RGBColor(0, 0, 0)
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2), Cm(2)
    section.left_margin, section.right_margin = Cm(3), Cm(1.5)

    normal = doc.styles["Normal"]
    _docx_fonts(normal, body_font)
    normal.font.size = Pt(size)
    normal.font.color.rgb = black
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for level in (1, 2, 3):
        style = doc.styles[f"Heading {level}"]
        _docx_fonts(style, body_font)
        style.font.size, style.font.bold = Pt(size), True
        style.font.italic = level == 3
        style.font.color.rgb = black
        style.paragraph_format.space_before = Pt(12 if level == 1 else 6)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    if title:
        heading = doc.add_paragraph()
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        heading.paragraph_format.space_before = Pt(6)
        heading.paragraph_format.space_after = Pt(12)
        heading.paragraph_format.keep_with_next = True
        run = heading.add_run(title)
        run.bold = True

    for block in parse_blocks(content):
        kind = block[0]
        if kind == "table":
            rows = _square(block[1])
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = True
            for r, row in enumerate(rows):
                _docx_row_flags(table.rows[r], header=r == 0)
                for c, value in enumerate(row):
                    paragraph = table.cell(r, c).paragraphs[0]
                    paragraph.paragraph_format.space_after = Pt(0)
                    if r == 0:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for chunk, bold, italic in inline_runs(value):
                        run = paragraph.add_run(chunk)
                        run.bold, run.italic, run.font.size = bold or r == 0, italic, Pt(13)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
            continue
        if kind == "heading":
            doc.add_heading(_plain(block[2]), block[1])
            continue
        if kind == "bullet":
            text, indent = ("+ " if block[1] else "- ") + block[2], Cm(1.5 if block[1] else 1)
        else:
            text, indent = block[2], Cm(1)
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.first_line_indent = indent
        for chunk, bold, italic in inline_runs(text):
            run = paragraph.add_run(chunk)
            run.bold, run.italic = bold, italic

    section.different_first_page_header_footer = True
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _docx_page_number(header)
    for run in header.runs:
        run.font.size = Pt(13)
    if title:
        doc.core_properties.title = title
    doc.save(str(path))


# ---------------------------------------------------------------- PDF


def _pdf_fonts() -> Tuple[str, str]:
    regular = os.environ.get("ZALO_PDF_FONT", "").strip()
    bold = os.environ.get("ZALO_PDF_FONT_BOLD", "").strip() or regular
    candidates = ([(regular, bold)] if regular else []) + list(_PDF_FONTS)
    for reg, bld in candidates:
        if reg and os.path.isfile(reg):
            return reg, bld if bld and os.path.isfile(bld) else reg
    raise FileSpecError("máy chủ thiếu font tiếng Việt để tạo PDF (đặt ZALO_PDF_FONT)")


def build_pdf(title: str, content: str, path: Path) -> None:
    from fpdf import FPDF
    from fpdf.fonts import FontFace

    regular, bold = _pdf_fonts()

    class Document(FPDF):
        def footer(self) -> None:
            self.set_y(-13)
            self.set_font("VN", "", 9)
            self.set_text_color(*_rgb(MUTED))
            self.cell(0, 8, f"Trang {self.page_no()}/{{nb}}", align="C")

    pdf = Document(format="A4")
    pdf.set_margins(20, 18, 20)
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_font("VN", "", regular)
    pdf.add_font("VN", "B", bold)
    fallback = next((f for f in _PDF_FALLBACK_FONTS if os.path.isfile(f) and f not in (regular, bold)), None)
    if fallback:
        pdf.add_font("VNSYM", "", fallback)
        pdf.set_fallback_fonts(["VNSYM"], exact_match=False)
    pdf.add_page()
    width = pdf.w - pdf.l_margin - pdf.r_margin

    def write(text: str, size: float, color: str = INK, style: str = "", indent: float = 0.0) -> None:
        pdf.set_font("VN", style, size)
        pdf.set_text_color(*_rgb(color))
        pdf.set_x(pdf.l_margin + indent)
        # fpdf2 hiểu **đậm**; nghiêng không có font riêng nên bỏ dấu *.
        safe = "".join(f"**{chunk}**" if b else chunk for chunk, b, _i in inline_runs(text))
        pdf.multi_cell(width - indent, size * 0.52, safe, markdown=not style, new_x="LMARGIN", new_y="NEXT")

    if title:
        write(title, 19, NAVY, "B")
        pdf.set_draw_color(*_rgb(ACCENT))
        pdf.set_line_width(0.8)
        y = pdf.get_y() + 1.5
        pdf.line(pdf.l_margin, y, pdf.w - pdf.r_margin, y)
        pdf.ln(6)
    for block in parse_blocks(content):
        kind = block[0]
        if kind == "table":
            pdf.ln(1)
            pdf.set_font("VN", "", 10.5)
            pdf.set_text_color(*_rgb(INK))
            pdf.set_draw_color(*_rgb(GRID))
            pdf.set_line_width(0.2)
            headings = FontFace(family="VN", emphasis="BOLD", color=(255, 255, 255), fill_color=_rgb(NAVY))
            with pdf.table(text_align="LEFT", headings_style=headings, cell_fill_color=_rgb(ZEBRA),
                           cell_fill_mode="ROWS", line_height=6.2, padding=1.6) as table:
                for row in _square(block[1]):
                    cells = table.row()
                    for value in row:
                        cells.cell(_plain(value))
            pdf.ln(3)
        elif kind == "heading":
            pdf.ln(2)
            write(_plain(block[2]), {1: 15, 2: 13.5, 3: 12.5}[block[1]], NAVY if block[1] == 1 else ACCENT, "B")
            pdf.ln(0.5)
        elif kind == "bullet":
            indent = 5.0 + 6.0 * block[1]
            pdf.set_font("VN", "B", 11.5)
            pdf.set_text_color(*_rgb(ACCENT))
            pdf.set_x(pdf.l_margin + indent)
            pdf.cell(5, 11.5 * 0.52, "•")
            pdf.set_font("VN", "", 11.5)
            pdf.set_text_color(*_rgb(INK))
            safe = "".join(f"**{chunk}**" if b else chunk for chunk, b, _i in inline_runs(block[2]))
            pdf.multi_cell(width - indent - 5, 11.5 * 0.52, safe, markdown=True, new_x="LMARGIN", new_y="NEXT")
            pdf.ln(0.8)
        else:
            write(block[2], 11.5)
            pdf.ln(1.5)
    if title:
        pdf.set_title(title)
    pdf.output(str(path))


# ---------------------------------------------------------------- PowerPoint


def build_pptx(title: str, slides: Sequence[Dict[str, Any]], path: Path) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt

    font = OFFICE_FONT
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    width, height = prs.slide_width, prs.slide_height
    blank = prs.slide_layouts[6]

    def rect(slide, x, y, w, h, color):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(color)
        shape.line.fill.background()
        shape.shadow.inherit = False
        return shape

    def text_frame(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
        frame = slide.shapes.add_textbox(x, y, w, h).text_frame
        frame.word_wrap = True
        frame.vertical_anchor = anchor
        return frame

    def add_runs(paragraph, text, size, color, bold=False):
        for chunk, strong, italic in inline_runs(text):
            run = paragraph.add_run()
            run.text = chunk
            run.font.name, run.font.size = font, Pt(size)
            run.font.bold, run.font.italic = bold or strong, italic
            run.font.color.rgb = RGBColor.from_string(color)

    def cover(text, subtitle=""):
        slide = prs.slides.add_slide(blank)
        rect(slide, 0, 0, width, height, NAVY)
        rect(slide, Inches(0.8), Inches(4.05), Inches(1.8), Inches(0.12), LIGHT)
        frame = text_frame(slide, Inches(0.8), Inches(1.4), width - Inches(1.6), Inches(2.5), MSO_ANCHOR.BOTTOM)
        add_runs(frame.paragraphs[0], text, 40, "FFFFFF", bold=True)
        if subtitle:
            sub = text_frame(slide, Inches(0.8), Inches(4.35), width - Inches(1.6), Inches(1.2))
            add_runs(sub.paragraphs[0], subtitle, 20, LIGHT)

    if title:
        cover(title)
    total = len(slides)
    for number, spec in enumerate(slides, 1):
        if not spec["bullets"]:
            cover(spec["title"])
            continue
        slide = prs.slides.add_slide(blank)
        rect(slide, 0, 0, width, Inches(1.25), NAVY)
        rect(slide, 0, Inches(1.25), width, Inches(0.06), ACCENT)
        heading = text_frame(slide, Inches(0.6), Inches(0.12), width - Inches(1.2), Inches(1.0), MSO_ANCHOR.MIDDLE)
        add_runs(heading.paragraphs[0], spec["title"], 28, "FFFFFF", bold=True)
        bullets = spec["bullets"]
        size = 24 if len(bullets) <= 5 else 20 if len(bullets) <= 8 else 16
        body = text_frame(slide, Inches(0.85), Inches(1.65), width - Inches(1.7), Inches(5.2))
        for index, item in enumerate(bullets):
            paragraph = body.paragraphs[0] if index == 0 else body.add_paragraph()
            paragraph.space_after = Pt(size * 0.55)
            add_runs(paragraph, "●  ", size * 0.7, ACCENT, bold=True)
            add_runs(paragraph, item, size, INK)
        footer = text_frame(slide, width - Inches(1.6), height - Inches(0.55), Inches(1.2), Inches(0.4))
        footer.paragraphs[0].alignment = PP_ALIGN.RIGHT
        add_runs(footer.paragraphs[0], f"{number}/{total}", 12, MUTED)
    if title:
        prs.core_properties.title = title
    prs.save(str(path))


# ---------------------------------------------------------------- Excel


def build_xlsx(title: str, sheets: Sequence[Dict[str, Any]], path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    header_fill = PatternFill("solid", fgColor=NAVY)
    zebra_fill = PatternFill("solid", fgColor=ZEBRA)
    side = Side(style="thin", color=GRID)
    border = Border(left=side, right=side, top=side, bottom=side)

    book = Workbook()
    book.remove(book.active)
    for spec in sheets:
        sheet = book.create_sheet(spec["name"])
        for row in spec["rows"]:
            sheet.append(row)
        columns = max(len(row) for row in spec["rows"])
        for r, row in enumerate(sheet.iter_rows(max_col=columns), 1):
            for cell in row:
                cell.border = border
                long_text = isinstance(cell.value, str) and len(cell.value) > 40
                if r == 1:
                    cell.font = Font(name=OFFICE_FONT, size=11, bold=True, color="FFFFFF")
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                else:
                    cell.font = Font(name=OFFICE_FONT, size=11, color=INK)
                    cell.alignment = Alignment(vertical="top", wrap_text=long_text)
                    if r % 2 == 0:
                        cell.fill = zebra_fill
        sheet.row_dimensions[1].height = 26
        for column in sheet.iter_cols(max_col=columns):
            longest = max((max(len(line) for line in str(c.value).splitlines() or [""])
                           for c in column if c.value is not None), default=8)
            sheet.column_dimensions[column[0].column_letter].width = min(50, max(10, longest + 3))
        sheet.freeze_panes = "A2"
        if len(spec["rows"]) > 1:
            sheet.auto_filter.ref = sheet.dimensions
        sheet.page_setup.orientation = "landscape" if columns > 6 else "portrait"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
        sheet.print_title_rows = "1:1"
    if title:
        book.properties.title = title
    book.save(str(path))


def make_file(
    fmt: str,
    title: str = "",
    *,
    directory: str,
    filename: Optional[str] = None,
    content: Optional[str] = None,
    slides: Any = None,
    sheets: Any = None,
) -> Path:
    """Kiểm tra giới hạn rồi dựng tệp trong ``directory``. Trả đường dẫn tệp đã tạo."""
    fmt = str(fmt or "").lower().lstrip(".")
    if fmt not in FORMATS:
        raise FileSpecError("chỉ tạo được tệp docx, pptx, xlsx hoặc pdf")
    title = str(title or "").strip()
    path = Path(directory) / safe_filename(filename or title, fmt)
    if fmt in ("docx", "pdf"):
        text = str(content or "")
        _check_text(title, text)
        (build_docx if fmt == "docx" else build_pdf)(title, text, path)
    elif fmt == "pptx":
        build_pptx(title[:MAX_TITLE_CHARS], _normalize_slides(slides), path)
    else:
        build_xlsx(title[:MAX_TITLE_CHARS], _normalize_sheets(sheets), path)
    return path
